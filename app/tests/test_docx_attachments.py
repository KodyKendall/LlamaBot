"""
Tests for .docx attachment support.

Docx files can't be sent to any LLM as raw base64, so the backend extracts
the document text server-side and injects it as a text content block. This
makes docx attachments work for every model, including text-only DeepSeek.
"""
import base64
import io

import pytest

from app.websocket.request_handler import (
    DOCX_MIME_TYPE,
    MAX_DOCX_CHARS,
    RequestHandler,
    extract_docx_text,
)


def make_docx_b64(paragraphs, table_rows=None):
    """Build a real .docx in memory and return it base64-encoded."""
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, cell_text in enumerate(row):
                table.cell(r, c).text = cell_text
    buf = io.BytesIO()
    doc.save(buf)
    return base64.b64encode(buf.getvalue()).decode("utf-8")


class TestExtractDocxText:
    def test_extracts_paragraphs(self):
        data = make_docx_b64(["Quarterly report", "Revenue grew 12%."])
        text = extract_docx_text(data)
        assert "Quarterly report" in text
        assert "Revenue grew 12%." in text

    def test_extracts_table_cells(self):
        data = make_docx_b64(["Summary"], table_rows=[["Region", "Sales"], ["EMEA", "42"]])
        text = extract_docx_text(data)
        assert "Region" in text
        assert "EMEA" in text
        assert "42" in text

    def test_truncates_huge_documents(self):
        data = make_docx_b64(["lorem ipsum " * 200] * 200)
        text = extract_docx_text(data)
        assert len(text) <= MAX_DOCX_CHARS + 200  # allow room for the truncation marker
        assert "[truncated]" in text

    def test_invalid_data_raises(self):
        with pytest.raises(Exception):
            extract_docx_text(base64.b64encode(b"not a docx").decode("utf-8"))


class TestBuildMessageContentWithDocx:
    def _build(self, attachments, llm_model="deepseek-v4-flash"):
        handler = RequestHandler(app=None)
        return handler._build_message_content({
            "message": "Please summarize the attached document.",
            "attachments": attachments,
            "llm_model": llm_model,
        })

    def test_docx_text_is_inlined_for_text_only_model(self):
        content = self._build([{
            "filename": "report.docx",
            "mime_type": DOCX_MIME_TYPE,
            "data": make_docx_b64(["Quarterly report", "Revenue grew 12%."]),
        }])
        assert isinstance(content, list)
        text_blocks = [b["text"] for b in content if b.get("type") == "text"]
        joined = "\n".join(text_blocks)
        assert "report.docx" in joined
        assert "Revenue grew 12%." in joined
        # The docx must NOT be reported as unsupported, and never sent as binary
        assert "doesn't support" not in joined
        assert all(b.get("type") != "file" for b in content)

    def test_corrupt_docx_does_not_crash_and_notes_failure(self):
        content = self._build([{
            "filename": "broken.docx",
            "mime_type": DOCX_MIME_TYPE,
            "data": base64.b64encode(b"garbage").decode("utf-8"),
        }])
        assert isinstance(content, list)
        joined = "\n".join(b["text"] for b in content if b.get("type") == "text")
        assert "broken.docx" in joined
        assert "could not be read" in joined


class TestUploadDocxToAssets:
    @pytest.mark.asyncio
    async def test_docx_upload_is_accepted(self, async_client, tmp_path, monkeypatch):
        import app.routers.api as api

        monkeypatch.setattr(api, "IMPORTS_DIR", str(tmp_path))
        payload = base64.b64decode(make_docx_b64(["hello"]))
        response = await async_client.post(
            "/api/upload-to-assets",
            files={"file": ("notes.docx", payload, DOCX_MIME_TYPE)},
        )
        assert response.status_code == 200
        body = response.json()
        assert body["filename"] == "notes.docx"
        assert (tmp_path / "notes.docx").exists()
