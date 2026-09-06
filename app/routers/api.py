"""API routes for LlamaBot."""

import json
import logging
import re
import os
import subprocess

from typing import Optional
from fastapi import APIRouter, Request, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import JSONResponse
from pydantic import BaseModel
from sqlmodel import Session

from app.models import User, ThreadMetadata
from app.dependencies import get_db_session, auth, admin_required, engineer_or_admin_required, get_current_user
from app.services.thread_service import get_thread_list
from app.services.user_service import (
    get_all_users, get_user_by_username, update_user, delete_user
)
from app.agents.leonardo.model_capabilities import get_model_capabilities
from app.agents.leonardo.openrouter_models import (
    API_KEY_ENV as OPENROUTER_API_KEY_ENV,
    get_openrouter_model,
    openrouter_models,
)
from app.agents.leonardo.model_policy import (
    enabled_default_model,
    is_model_enabled,
    model_switching_allowed,
    policy_report,
    vision_allowed,
    vision_model,
)

logger = logging.getLogger(__name__)

router = APIRouter()


# ============== Pydantic Models ==============

class CreateUserRequest(BaseModel):
    username: str
    password: str
    is_admin: bool = False
    role: str = "engineer"


class UpdateUserRequest(BaseModel):
    is_active: bool | None = None
    is_admin: bool | None = None
    new_password: str | None = None
    role: str | None = None


class CreatePromptRequest(BaseModel):
    name: str
    content: str
    group: str = "General"
    description: str | None = None


class UpdatePromptRequest(BaseModel):
    name: str | None = None
    content: str | None = None
    group: str | None = None
    description: str | None = None
    is_active: bool | None = None


class WriteSkillRequest(BaseModel):
    name: str
    description: str = ""
    content: str
    slug: str | None = None


# ============== Version API ==============

def get_container_version() -> str:
    """Get the version from docker-compose.yml (mounted at /app/leonardo) or Docker API."""
    import socket

    # First, try to parse from mounted docker-compose.yml
    compose_paths = [
        "/app/leonardo/docker-compose.yml",
        "/app/leonardo/docker-compose-dev.yml",
    ]
    for compose_path in compose_paths:
        try:
            with open(compose_path, 'r') as f:
                for line in f:
                    # Look for image line with llamabot (commented or not)
                    # e.g., "# image: kody06/llamabot:0.3.5c" or "image: kody06/llamabot:0.3.5c"
                    if 'image:' in line and 'llamabot:' in line:
                        # Extract version from "kody06/llamabot:0.3.5c"
                        match = re.search(r'llamabot:([^\s"\']+)', line)
                        if match:
                            return match.group(1)
        except FileNotFoundError:
            continue
        except Exception as e:
            logger.debug(f"Could not parse {compose_path}: {e}")

    # Fallback: Query Docker API
    try:
        container_id = socket.gethostname()
        import http.client
        conn = http.client.HTTPConnection("localhost")
        conn.sock = socket.socket(socket.AF_UNIX, socket.SOCK_STREAM)
        conn.sock.connect("/var/run/docker.sock")
        conn.request("GET", f"/containers/{container_id}/json")
        response = conn.getresponse()

        if response.status == 200:
            data = json.loads(response.read().decode())
            image = data.get("Config", {}).get("Image", "")
            if ":" in image:
                version = image.split(":")[-1]
                if version != "latest":
                    return version
        return "dev"
    except Exception as e:
        logger.debug(f"Could not get container version: {e}")
        return "dev"


def get_llamapress_version() -> str:
    """Get the LlamaPress version from docker-compose.yml."""
    compose_paths = [
        "/app/leonardo/docker-compose.yml",
        "/app/leonardo/docker-compose-dev.yml",
    ]
    for compose_path in compose_paths:
        try:
            with open(compose_path, 'r') as f:
                for line in f:
                    if 'image:' in line and 'llamapress-simple:' in line:
                        match = re.search(r'llamapress-simple:([^\s"\']+)', line)
                        if match:
                            return match.group(1)
        except FileNotFoundError:
            continue
        except Exception as e:
            logger.debug(f"Could not parse {compose_path} for llamapress version: {e}")
    return "dev"


@router.get("/api/version", response_class=JSONResponse)
async def api_get_version():
    """Get the current LlamaBot version from docker-compose.yml or Docker image tag."""
    version = get_container_version()
    return {"version": version}


def get_version_notes(version: str) -> list[str]:
    """
    Get release notes for a specific version from docs/dev_logs.

    Args:
        version: Full version string like "0.3.6h"

    Returns:
        List of note strings for that version, or empty list if not found
    """
    if not version or version == "dev":
        return []

    # Parse version to get base (e.g., "0.3.6" from "0.3.6h")
    # Version format: X.Y.Z or X.Y.Za where 'a' is optional letter suffix
    match = re.match(r'^(\d+\.\d+\.\d+)', version)
    if not match:
        return []

    base_version = match.group(1)

    # Try to read the dev_log file
    dev_log_paths = [
        f"/app/docs/dev_logs/{base_version}",
        f"/app/app/docs/dev_logs/{base_version}",  # Alternative path in container
    ]

    content = None
    for path in dev_log_paths:
        try:
            with open(path, 'r') as f:
                content = f.read()
            break
        except FileNotFoundError:
            continue
        except Exception as e:
            logger.debug(f"Could not read dev_log {path}: {e}")
            continue

    if not content:
        return []

    # Find the section for this version (e.g., "0.3.6h:" or "0.3.6ec:")
    # Pattern: version at start of line followed by colon
    lines = content.split('\n')
    notes = []
    in_section = False

    for line in lines:
        stripped = line.strip()

        # Check if this is a version header (e.g., "0.3.6h:" or "0.3.6ec:")
        if re.match(r'^\d+\.\d+\.\d+[a-z]*:', stripped):
            if in_section:
                # We've reached the next version, stop collecting
                break
            # Check if this is our version
            if stripped.startswith(f"{version}:"):
                in_section = True
            continue

        # Collect bullet points while in our section
        if in_section and stripped.startswith('- '):
            # Clean up the note: remove "- [x] " or "- [ ] " prefix
            note = re.sub(r'^- \[[x ]\] ', '', stripped)
            note = re.sub(r'^- ', '', note)  # Also handle plain "- " prefix
            if note:
                notes.append(note)

    return notes


@router.get("/api/version-notes", response_class=JSONResponse)
async def api_get_version_notes():
    """Get the current version and its release notes from docs/dev_logs."""
    version = get_container_version()
    notes = get_version_notes(version)
    return {"version": version, "notes": notes}


# A per-version-pair "this update already failed here" marker, stored in
# SiteSetting. Keeps the banner from nagging (and re-burning a dead wait) after
# a failed attempt; the next release is a new pair, so it un-sticks itself.
# Deliberately NOT in VALID_SITE_SETTINGS — private, like session_secret.

def _update_failure_key(lb_version: str, lp_version: str) -> Optional[str]:
    key = f"update_failed:{lb_version}:{lp_version}"
    return key if len(key) <= 100 else None  # SiteSetting.key max_length


def _get_update_failure(session: Session, lb_version: str, lp_version: str) -> Optional[str]:
    key = _update_failure_key(lb_version, lp_version)
    if key is None:
        return None
    return get_site_setting(session, key, default="") or None


def _record_update_failure(session: Session, lb_version: str, lp_version: str, error_tail: str) -> None:
    """Best-effort: a down auth DB must never mask the honest failure response."""
    try:
        from datetime import datetime, timezone
        from app.models import SiteSetting

        key = _update_failure_key(lb_version, lp_version)
        if key is None:
            return
        value = json.dumps({
            "at": datetime.now(timezone.utc).isoformat(),
            "error": error_tail[:900],  # SiteSetting.value max_length=1000
        })
        setting = session.get(SiteSetting, key)
        if setting:
            setting.value = value
            setting.updated_at = datetime.now(timezone.utc)
        else:
            session.add(SiteSetting(key=key, value=value))
        session.commit()
    except Exception as e:
        logger.warning(f"Could not record update failure marker: {e}")


async def _report_update_failure(lb_version: str, lp_version: str,
                                 error_class: str, error_tail: str, detail: str) -> None:
    """Fire-and-forget mothership telemetry; never interferes with the response."""
    try:
        from app.services.mothership_client import MothershipClient

        await MothershipClient().report_error(
            thread_id=None,
            error_class=error_class,
            error_message=error_tail,
            traceback_str=detail,
            fingerprint=f"update_failed:{lb_version}:{lp_version}",
            llamabot_version=get_container_version(),
            recovered=False,
        )
    except Exception as e:
        logger.warning(f"Could not report update failure to mothership: {e}")


@router.get("/api/check-updates", response_class=JSONResponse)
async def api_check_updates(session: Session = Depends(get_db_session)):
    """Check mothership for available updates to llamabot and llamapress."""
    from app.services.mothership_client import MothershipClient
    mothership = MothershipClient()
    if not mothership.enabled:
        return {"updates_available": False, "reason": "mothership_not_configured"}

    llamabot_version = get_container_version()
    llamapress_version = get_llamapress_version()

    result = await mothership.check_updates(llamabot_version, llamapress_version)
    if result is None:
        return {"updates_available": False, "reason": "check_failed"}

    # Don't re-offer a pair that already failed on this instance — the user
    # clicked, it broke, and mothership was told. Fails open on DB trouble.
    if result.get("updates_available"):
        latest = result.get("latest_versions") or {}
        target_lb = (latest.get("llamabot") or {}).get("version") or llamabot_version
        target_lp = (latest.get("llamapress") or {}).get("version") or llamapress_version
        if _get_update_failure(session, target_lb, target_lp):
            return {"updates_available": False, "reason": "previous_update_failed"}
    return result


@router.get("/api/overlay-ads", response_class=JSONResponse)
async def api_overlay_ads(request: Request):
    """Promo snippets AND the display policy for the building overlay.

    Thin, fail-open proxy in front of the mothership: it owns the creative, the
    cadence, and the rules for when a promo shows at all; we only clamp and
    cache (see ``app.services.overlay_ads``). An unconfigured box, an old
    mothership, or any error all return an empty list with default policy — the
    overlay then shows no slot, exactly as before this shipped.

    The signed-in user is passed along so the mothership can personalise both
    halves. Resolved *optionally*: this must never 401, because a promo failing
    to load is not a reason to break the build overlay.
    """
    from app.dependencies import _user_from_session_cookie
    from app.services import overlay_ads
    from app.services.mothership_client import MothershipClient

    # A local override file is an explicit operator choice, so it wins over both
    # the cache and the mothership — edit the file, reload, see the new promo.
    local = overlay_ads.local_override()
    if local is not None:
        return local

    # Resolved WITHOUT a FastAPI dependency on purpose: `Depends(get_db_session)`
    # would make a DB hiccup 500 an endpoint whose whole contract is to fail open.
    # No user just means no personalisation.
    user = None
    try:
        from app.db import engine
        if engine is not None:
            with Session(engine) as db:
                user = _user_from_session_cookie(request, db)
    except Exception as e:
        logger.info(f"Overlay ads: could not resolve user, serving unpersonalised: {e}")

    # Cache per user: the mothership personalises this response, so a shared
    # cache would serve one user's targeted promos and policy to the next.
    cache_key = f"user:{user.id}" if user else "anon"

    fresh = overlay_ads.cached(cache_key)
    if fresh is not None:
        return fresh

    mothership = MothershipClient()
    if not mothership.enabled:
        return overlay_ads.empty()

    # Same wire shape every mothership report uses, so a promo can be targeted
    # by the same llamapress_user_guid the telemetry is keyed on.
    from app.services import user_context as user_ctx

    raw = await mothership.fetch_overlay_ads(
        get_container_version(), user=user_ctx.describe(user)
    )
    if raw is None:
        # Cache the miss too, so an unreachable mothership isn't re-dialed on
        # every single build for the next minute.
        return overlay_ads.store(overlay_ads.empty(), cache_key)
    return overlay_ads.store(overlay_ads.normalize(raw), cache_key)


@router.get("/api/rails-errors", response_class=JSONResponse)
async def api_rails_errors(request: Request, username: str = Depends(auth)):
    """The Rails app's recent crashes, for the chat page's error tray.

    The tray above the composer already shows JavaScript errors the preview
    pushes over postMessage. This is the other half: press a button, get a 500,
    and the same notice says so — instead of the user staring at a Rails error
    page wondering whether Leo can see it.

    Why LlamaBot proxies instead of the browser reading Rails directly:
    ``GET /llama_bot/errors`` sets no CORS headers and the chat page is a
    different origin, so a direct fetch is blocked. Proxying also keeps the
    whole feature inside the LlamaBot image — it works against any gem already
    serving the feed, with no skeleton release in the way.

    Polled every few seconds by a signed-in user, so it stays cheap and quiet:
    the work is one request on the Docker network, and every failure answers
    "no information" rather than something the tray would have to render.
    """
    return await _rails_errors(request)


def _rails_errors_feed(token: str):
    from app.services.rails_error_feed import RailsErrorFeedClient

    return RailsErrorFeedClient(token=token)


async def _rails_errors(request: Request, feed_factory=_rails_errors_feed):
    """Everything the route does. Split out so tests can inject a fake feed."""
    from app.lib import rails_error_tray

    unavailable = {"seq": None, "errors": [], "available": False}

    # Header, never a query parameter: this is a live 30-minute Rails bearer
    # token and query strings end up in access logs and browser history.
    token = (request.headers.get("X-Rails-Api-Token") or "").strip()
    if not token:
        # Signed into LlamaBot but not into the Rails app. Nothing to poll.
        return unavailable

    since = _rails_error_cursor(request.query_params.get("since"))

    try:
        result = await feed_factory(token).fetch(since=since)
    except Exception as e:  # noqa: BLE001 - a broken feed is not a broken page
        logger.debug(f"Rails error feed unavailable: {e}")
        return unavailable

    if result is None:
        # Gem too old for the endpoint, token expired, or Rails restarting.
        return unavailable

    seq, errors = result
    return {
        "seq": seq,
        "errors": rails_error_tray.tray_entries(errors),
        "available": True,
    }


def _rails_error_cursor(raw) -> Optional[int]:
    """The caller's cursor, or None to make this a probe.

    Anything that is not a plain non-negative integer becomes a probe rather
    than 0 — reading a garbled cursor as 0 would replay the whole ring into the
    tray, which is how the user would get shown crashes from yesterday.
    """
    value = str(raw or "").strip()
    if not value.isdigit():
        return None
    return int(value)


class UpdateRequest(BaseModel):
    llamabot_version: str
    llamapress_version: str


@router.post("/api/update", response_class=JSONResponse)
async def api_perform_update(
    request: UpdateRequest,
    current_user: User = Depends(engineer_or_admin_required),
    session: Session = Depends(get_db_session),
):
    """Update docker-compose.yml image tags, pull new images, and restart.

    Failures are reported honestly: bin/update exiting non-zero, a timeout, or
    a broken host bridge all return success=false with an error tail, persist a
    per-pair failure marker (so /api/check-updates stops offering that pair),
    and tell the mothership. A container killed mid-restart produces no
    response at all — the frontend treats a dropped connection as the restart,
    so nothing catchable here is a success.
    """
    version_pattern = re.compile(r'^[0-9a-zA-Z.\-]+$')
    if not version_pattern.match(request.llamabot_version) or not version_pattern.match(request.llamapress_version):
        raise HTTPException(status_code=400, detail="Invalid version format")

    from app.routers.slash_commands import execute_command
    command = f"bash bin/update {request.llamabot_version} {request.llamapress_version}"

    async def _fail(error_class: str, error_tail: str, detail: str, return_code: int):
        _record_update_failure(session, request.llamabot_version, request.llamapress_version, error_tail)
        await _report_update_failure(request.llamabot_version, request.llamapress_version,
                                     error_class, error_tail, detail)
        return {
            "success": False,
            "stdout": "",
            "stderr": "",
            "return_code": return_code,
            "error": error_tail,
        }

    try:
        result = execute_command(command, timeout=300)
    except subprocess.TimeoutExpired as e:
        return await _fail("UpdateTimeout", f"Update timed out after {e.timeout}s", str(e), -1)
    except Exception as e:
        return await _fail(type(e).__name__, str(e), str(e), -1)

    if result.returncode != 0:
        output = (result.stderr or "").strip() or (result.stdout or "").strip() \
            or f"bin/update exited {result.returncode}"
        return await _fail("UpdateFailed", output[-500:],
                           f"stdout:\n{result.stdout}\n\nstderr:\n{result.stderr}",
                           result.returncode)

    return {
        "success": True,
        "stdout": result.stdout.strip() if result.stdout else "",
        "stderr": result.stderr.strip() if result.stderr else "",
        "return_code": 0,
    }


# ============== WebSocket Authentication API ==============

@router.get("/api/ws-token", response_class=JSONResponse)
async def get_ws_token(current_user: User = Depends(get_current_user)):
    """
    Generate a JWT token for WebSocket authentication.

    The token is used by the frontend to authenticate WebSocket connections.
    It expires after WS_TOKEN_EXPIRY_MINUTES (default 30 minutes).
    """
    from app.services.token_service import create_ws_token, EXPIRY_MINUTES
    token = create_ws_token(current_user)
    return {"token": token, "expires_in": EXPIRY_MINUTES * 60}


# ============== User Management API ==============

@router.get("/api/users", response_class=JSONResponse)
async def api_get_users(
    admin: User = Depends(admin_required),
    session: Session = Depends(get_db_session)
):
    """Get all users (admin only)."""
    users = get_all_users(session)
    return [
        {
            "id": u.id,
            "username": u.username,
            "is_active": u.is_active,
            "is_admin": u.is_admin,
            "role": getattr(u, 'role', 'engineer'),
            "created_at": u.created_at.isoformat() if u.created_at else None,
            "updated_at": u.updated_at.isoformat() if u.updated_at else None,
        }
        for u in users
    ]


@router.post("/api/users", response_class=JSONResponse)
async def api_create_user(
    request: CreateUserRequest,
    admin: User = Depends(admin_required),
    session: Session = Depends(get_db_session)
):
    """Create a new user (admin only)."""
    from app.services.user_service import hash_password, sanitize_username

    # Sanitize username (strip whitespace)
    try:
        clean_username = sanitize_username(request.username)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    # Check if username exists
    if get_user_by_username(session, clean_username):
        raise HTTPException(status_code=409, detail="Username already exists")

    user = User(
        username=clean_username,
        password_hash=hash_password(request.password),
        is_admin=request.is_admin,
        role=request.role
    )
    session.add(user)
    session.commit()
    session.refresh(user)

    logger.info(f"Admin '{admin.username}' created user '{clean_username}' with role '{request.role}'")
    return {"id": user.id, "username": user.username, "message": "User created successfully"}


@router.patch("/api/users/{user_id}", response_class=JSONResponse)
async def api_update_user(
    user_id: int,
    request: UpdateUserRequest,
    admin: User = Depends(admin_required),
    session: Session = Depends(get_db_session)
):
    """Update a user (admin only)."""
    # Prevent admin from deactivating themselves
    if request.is_active is False and user_id == admin.id:
        raise HTTPException(status_code=400, detail="Cannot deactivate yourself")

    # Prevent admin from removing their own admin status
    if request.is_admin is False and user_id == admin.id:
        raise HTTPException(status_code=400, detail="Cannot remove your own admin status")

    user = update_user(
        session, user_id,
        is_active=request.is_active,
        is_admin=request.is_admin,
        new_password=request.new_password,
        role=request.role
    )

    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    logger.info(f"Admin '{admin.username}' updated user {user_id}")
    return {"message": "User updated successfully"}


@router.delete("/api/users/{user_id}", response_class=JSONResponse)
async def api_delete_user(
    user_id: int,
    admin: User = Depends(admin_required),
    session: Session = Depends(get_db_session)
):
    """Delete a user (admin only)."""
    # Prevent admin from deleting themselves
    if user_id == admin.id:
        raise HTTPException(status_code=400, detail="Cannot delete yourself")

    if not delete_user(session, user_id):
        raise HTTPException(status_code=404, detail="User not found")

    logger.info(f"Admin '{admin.username}' deleted user {user_id}")
    return {"message": "User deleted successfully"}


# ============== Thread/Chat History API ==============

@router.get("/threads", response_class=JSONResponse)
async def threads(
    request: Request,
    username: str = Depends(auth),
    before: Optional[str] = Query(None, description="Cursor for pagination - ISO timestamp"),
    limit: int = Query(10, ge=1, le=50)
):
    """Get recent conversation threads with cursor-based pagination (fast - metadata only).

    This endpoint queries the lightweight ThreadMetadata table instead of loading
    full LangGraph checkpoint states, providing a massive performance improvement.
    """
    from datetime import datetime
    from app.db import engine

    if engine is None:
        logger.error("Database engine not available")
        return {"threads": [], "next_cursor": None, "has_more": False}

    with Session(engine) as session:
        # Parse cursor timestamp
        before_dt = None
        if before:
            try:
                before_dt = datetime.fromisoformat(before.replace('Z', '+00:00'))
            except ValueError:
                logger.warning(f"Invalid cursor timestamp: {before}")

        # Query lightweight metadata (no checkpoint loading!)
        threads = get_thread_list(session, before=before_dt, limit=limit + 1)

        # Check if there are more results
        has_more = len(threads) > limit
        threads = threads[:limit]

        # Build response
        next_cursor = threads[-1].updated_at.isoformat() if threads else None

        logger.info(f"Returning {len(threads)} threads from metadata table (fast query)")

        return {
            "threads": [
                {
                    "thread_id": t.thread_id,
                    "title": t.title,
                    "created_at": t.created_at.isoformat() + "Z",
                    "updated_at": t.updated_at.isoformat() + "Z",
                    "message_count": t.message_count,
                    "agent_name": t.agent_name
                }
                for t in threads
            ],
            "next_cursor": next_cursor,
            "has_more": has_more
        }


class CompactThreadRequest(BaseModel):
    thread_id: str
    agent_name: str = "rails_agent"


@router.post("/api/compact-thread", response_class=JSONResponse)
async def compact_thread(request_data: CompactThreadRequest, request: Request, username: str = Depends(auth)):
    """Compact the conversation history for a thread by summarizing older messages.

    Mirrors what SummarizationMiddleware does automatically at the token threshold,
    but runs on-demand so the user can trigger it via the /compact slash command.
    Keeps the last 15 messages verbatim and replaces everything before them with a
    structured summary generated by the same summarization model.
    """
    from langchain.agents.middleware import SummarizationMiddleware
    from langchain_core.messages import RemoveMessage
    from langgraph.graph.message import REMOVE_ALL_MESSAGES
    from app.agents.leonardo.llm_factory import make_summarization_model
    from app.agents.leonardo.rails_agent.nodes import SUMMARIZATION_PROMPT

    app = request.app
    agent_name = request_data.agent_name

    graph = app.state.compiled_graphs.get(agent_name)
    if not graph:
        # All agents share the same MessagesState schema + checkpointer, so any
        # compiled graph can read/write a thread. Fall back to rails_agent → llamabot.
        graph = app.state.compiled_graphs.get("rails_agent") or app.state.compiled_graphs.get("llamabot")
        if not graph:
            raise HTTPException(
                status_code=503,
                detail="No compiled graph available — the server may still be starting up. Try again in a moment."
            )
        logger.info(f"/compact: agent '{agent_name}' not in compiled_graphs, using fallback graph")

    config = {"configurable": {"thread_id": request_data.thread_id}}
    state_snapshot = await graph.aget_state(config=config)

    if not state_snapshot or not state_snapshot.values:
        raise HTTPException(status_code=404, detail="Thread not found or has no messages")

    messages = list(state_snapshot.values.get("messages", []))
    keep_count = 15

    if len(messages) <= keep_count:
        return {
            "success": True,
            "message": f"Only {len(messages)} messages — nothing to compact (need more than {keep_count})",
            "summarized": 0,
            "preserved": len(messages),
        }

    model, token_counter, trim_tokens_to_summarize = make_summarization_model()

    middleware = SummarizationMiddleware(
        model=model,
        trigger=("tokens", 1),  # threshold of 1 ensures it always fires
        keep=("messages", keep_count),
        token_counter=token_counter,
        trim_tokens_to_summarize=trim_tokens_to_summarize,
        summary_prompt=SUMMARIZATION_PROMPT,
    )

    middleware._ensure_message_ids(messages)
    cutoff_index = middleware._determine_cutoff_index(messages)

    if cutoff_index <= 0:
        return {
            "success": True,
            "message": "No messages need to be summarized",
            "summarized": 0,
            "preserved": len(messages),
        }

    messages_to_summarize, preserved_messages = middleware._partition_messages(messages, cutoff_index)
    summary = await middleware._acreate_summary(messages_to_summarize)
    new_messages = middleware._build_new_messages(summary)

    update = {
        "messages": [
            RemoveMessage(id=REMOVE_ALL_MESSAGES),
            *new_messages,
            *preserved_messages,
        ]
    }

    await graph.aupdate_state(config, update)

    logger.info(
        f"User '{username}' compacted thread '{request_data.thread_id}': "
        f"summarized {len(messages_to_summarize)} messages, preserved {len(preserved_messages)}"
    )

    return {
        "success": True,
        "message": (
            f"Compacted {len(messages_to_summarize)} messages into a summary; "
            f"kept {len(preserved_messages)} recent messages verbatim"
        ),
        "summarized": len(messages_to_summarize),
        "preserved": len(preserved_messages),
        "total_before": len(messages),
        "total_after": 1 + len(preserved_messages),
    }


@router.get("/api/thread-tokens/{thread_id}", response_class=JSONResponse)
async def thread_token_count(
    thread_id: str,
    request: Request,
    agent_name: str = Query(default="rails_agent"),
    username: str = Depends(auth),
):
    """Return the tiktoken count of the current checkpoint for a thread.

    Uses the same text-only token counter as SummarizationMiddleware so the
    result matches the threshold that triggers auto-compact.  Counts all message
    content — tool call arguments, tool results, reasoning blocks, etc. — which
    the old client-side character estimate silently missed.  No external API call:
    tiktoken is local and synchronous, so this endpoint is cheap to call on every
    thread load.
    """
    from app.agents.utils.token_counter import tiktoken_token_counter

    app = request.app
    graph = (
        app.state.compiled_graphs.get(agent_name)
        or app.state.compiled_graphs.get("rails_agent")
        or app.state.compiled_graphs.get("llamabot")
    )
    if not graph:
        return {"token_count": 0, "message_count": 0}

    config = {"configurable": {"thread_id": thread_id}}
    try:
        state_snapshot = await graph.aget_state(config=config)
    except Exception:
        return {"token_count": 0, "message_count": 0}

    if not state_snapshot or not state_snapshot.values:
        return {"token_count": 0, "message_count": 0}

    messages = list(state_snapshot.values.get("messages", []))
    token_count = tiktoken_token_counter(messages) if messages else 0

    return {"token_count": token_count, "message_count": len(messages)}


@router.get("/chat-history/{thread_id}")
async def chat_history(thread_id: str, request: Request, username: str = Depends(auth)):
    """Get chat history for a specific thread."""
    app = request.app
    checkpointer = app.state.get_or_create_checkpointer()

    # Use cached graph from startup (singleton pattern)
    graph = app.state.compiled_graphs.get("llamabot")
    if not graph:
        from app.agents.llamabot.nodes import build_workflow
        graph = build_workflow(checkpointer=checkpointer)
        logger.warning("/chat-history endpoint using fallback graph compilation")

    config = {"configurable": {"thread_id": thread_id}}
    state_history = await graph.aget_state(config=config)
    print(state_history)
    return state_history


# ============== Other API Endpoints ==============

@router.get("/available-agents", response_class=JSONResponse)
async def available_agents():
    """Get list of available agents (platform base ∪ client overlay)."""
    from app.lib.langgraph_registry import load_graphs
    return {"agents": list(load_graphs().keys())}


@router.get("/api/available-models", response_class=JSONResponse)
async def available_models(request: Request):
    """Get list of available LLM models based on configured API keys.

    Returns which models are available (have API keys) and which are not.
    Frontend uses this to disable unavailable models in the dropdown.
    """
    # Map of model frontend values to their required API key env vars
    # Values can be a single string or tuple of strings (checked in order, first found wins)
    model_api_keys = {
        "claude-4.5-haiku": "ANTHROPIC_API_KEY",
        "claude-4.5-sonnet": "ANTHROPIC_API_KEY",
        "gpt-5-mini": "OPENAI_API_KEY",
        "gpt-5-codex": "OPENAI_API_KEY",
        "gpt-5-nano": "OPENAI_API_KEY",
        "gpt-5.4-nano": "OPENAI_API_KEY",
        "gpt-5.6-luna": "OPENAI_API_KEY",
        "gemini-3-flash": ("GEMINI_API_KEY", "GOOGLE_API_KEY"),
        "gemini-3-pro": ("GEMINI_API_KEY", "GOOGLE_API_KEY"),
        "gemini-3.1-flash-lite": ("GEMINI_API_KEY", "GOOGLE_API_KEY"),
        "deepseek-v4-flash": "DEEPSEEK_API_KEY",
        "deepseek-v4-pro": "DEEPSEEK_API_KEY",
        # DeepSeek's vision sibling — same key as the text models, which is the
        # whole point of it (vision with no extra credential to provision).
        "deepseek-v4-flash-vision-exp": "DEEPSEEK_API_KEY",
        "deepseek-v4-flash-gmi": "GMI_DEEPSEEK_API_KEY",
        "deepseek-v4-flash-fireworks": "FIREWORKS_DEEPSEEK_API_KEY",
        # Fireworks' account-wide key name, falling back to the DeepSeek-specific
        # name already deployed on boxes — matches get_llm's precedence.
        "nemotron-lightning-30b-fireworks": ("FIREWORKS_API_KEY", "FIREWORKS_DEEPSEEK_API_KEY"),
        "qwen3.7-plus": "ALIBABA_API_KEY",
        "qwen3.8-27b-hetzner": "HETZNER_API_KEY",
        # Self-hosted vLLM pod: the "key" here is deliberately the BASE URL, not
        # an API key. Availability means "this box is pointed at a RunPod
        # endpoint" — the pod may legitimately run unauthenticated, so keying it
        # on RUNPOD_QWEN_API_KEY would grey out a working model everywhere.
        "qwen3-8b-runpod": "RUNPOD_QWEN_BASE_URL",
        "muse-glimmer-30b-runpod": "RUNPOD_GLIMMER_BASE_URL",
        "nemotron-lightning-30b-runpod": "RUNPOD_NEMOTRON_BASE_URL",
        # Meta's docs call the key MODEL_API_KEY; their LiteLLM integration calls
        # it META_API_KEY. Accept either, matching get_llm's precedence.
        "muse-spark-1.2-contributor": ("META_API_KEY", "MODEL_API_KEY"),
        # 1.3 is the same account and the same key as 1.2 — one Meta Model API
        # key covers both ids, so a per-model key name would be fiction.
        "muse-spark-1.3-contributor": ("META_API_KEY", "MODEL_API_KEY"),
    }

    # Models paid for by the SIGNED-IN USER's ChatGPT plan. Their availability is
    # not an env var — it is whether this user has connected their account — so
    # they are resolved separately below and skipped by the env-var loop.
    from app.agents.leonardo.llm_factory import _CHATGPT_SUBSCRIPTION_MODELS

    chatgpt_connected = False
    try:
        from sqlmodel import Session

        from app.db import engine
        from app.dependencies import try_authenticate
        from app.services.chatgpt_auth import status_for_user

        if engine is not None:
            with Session(engine) as db:
                current = try_authenticate(request, db)
                if current is not None:
                    chatgpt_connected = bool(
                        status_for_user(db, current.id).get("connected")
                    )
    except Exception as e:
        # Dropdown shaping must never 500 the chat page.
        logger.warning("Could not resolve ChatGPT connection status: %s", e)

    # Config-driven registry entries (see openrouter_models). They are not in
    # model_api_keys above — that map is hand-maintained per model, and the whole
    # point of the registry is that adding an endpoint touches no code — so they
    # are folded in here. Since 0.7.7 each entry carries its OWN key env(s)
    # rather than all sharing OPENROUTER_API_KEY: an entry may name any
    # OpenAI-compatible gateway, and greying it out on the wrong key would say
    # "not configured" about a model the box can perfectly well run.
    for model_value, registry_entry in openrouter_models().items():
        model_api_keys.setdefault(model_value, tuple(registry_entry["api_key_env"]))

    models = []
    for model_value, env_vars in model_api_keys.items():
        # Support both single string and tuple of env vars
        if isinstance(env_vars, str):
            env_vars = (env_vars,)

        # Check each env var in order, use first one that has a value
        has_key = False
        for env_var in env_vars:
            api_key = os.environ.get(env_var, "")
            if api_key and api_key.strip():
                has_key = True
                break

        # Operator/mothership policy can disable a model regardless of its key
        # (see model_policy). A disabled model is greyed out in the dropdown; the
        # real enforcement is in get_llm, since this endpoint is UX only.
        enabled = is_model_enabled(model_value)
        if not enabled:
            reason = "Disabled by administrator"
        elif not has_key:
            # Deliberately does NOT name the env var. This endpoint is readable by
            # every signed-in user, and the variable name is a fact about the
            # instance's .env — which nothing user-facing discloses. "Which key is
            # missing" is an operator question, answered by the server logs.
            reason = "API key not configured"
        else:
            reason = None

        entry = {
            "value": model_value,
            "available": has_key and enabled,
            "reason": reason,
            "capabilities": get_model_capabilities(model_value),
        }
        # Registry models have no <option> in chat.html — they are config, so the
        # markup cannot know their names. Ship their labels and the frontend
        # creates the option. Only sent for those: for a hardcoded model the
        # dropdown's own markup stays authoritative.
        registry_entry = get_openrouter_model(model_value)
        if registry_entry is not None:
            entry["label"] = registry_entry["label"]
            entry["short_label"] = registry_entry["short_label"]
        models.append(entry)

    for model_value in _CHATGPT_SUBSCRIPTION_MODELS:
        enabled = is_model_enabled(model_value)
        if not enabled:
            reason = "Disabled by administrator"
        elif not chatgpt_connected:
            reason = "Connect your ChatGPT account to use this model"
        else:
            reason = None
        models.append({
            "value": model_value,
            "available": chatgpt_connected and enabled,
            "reason": reason,
            "capabilities": get_model_capabilities(model_value),
            "requires_chatgpt_login": True,
        })

    # Coarse operator gates the frontend needs to shape the UI: hide the model
    # dropdown when switching is locked, and refuse image attachments (with a
    # support message) when vision is off. Both are re-enforced server-side
    # (get_llm / _build_message_content); these flags are UX only.
    return {
        "models": models,
        "model_switching_allowed": model_switching_allowed(),
        "vision_allowed": vision_allowed(),
        # Which model the frontend resets to on a new thread and pins to under
        # the switching lock. Box-dependent since 0.7.0 (Muse where there is a
        # META key, DeepSeek where there is not), so it can no longer be a
        # constant in index.js.
        "default_model": enabled_default_model(),
        # Where the image auto-switch sends a user who attaches an image while on
        # a text-only model. Box-dependent for the same reason since 0.7.5 (Muse
        # where there is a META key, DeepSeek's vision model where there is not),
        # so index.js can no longer hardcode Muse. Empty string means this box has
        # no vision model at all, which is what turns on the "no image-capable
        # model is configured" banner.
        "vision_model": vision_model(),
        # The resolved routing AND which channel decided it (0.7.7). Not used by
        # the frontend — this is here so "what is this box actually running, and
        # who told it that" is answerable over HTTP instead of over SSH, which is
        # what made the 2026-08-31 incident a three-hour one. See
        # model_policy.policy_report for why it reports the resolved value rather
        # than the configured intent.
        "policy": policy_report(),
    }


@router.get("/rails-routes", response_class=JSONResponse)
async def rails_routes():
    """Parse routes.rb and return available GET routes for `index` actions and home page."""
    routes = []
    routes_file = "rails/config/routes.rb"

    # Check if routes file exists
    if not os.path.exists(routes_file):
        return {"routes": [{"path": "/", "name": "Home"}]}

    try:
        with open(routes_file, "r") as f:
            content = f.read()

        # Extract root path
        root_match = re.search(r'root\s+"([^"]+)#([^"]+)"', content)
        if root_match:
            routes.append({"path": "/", "name": "Home"})

        # Extract explicit home route
        home_match = re.search(r'get\s+"home"\s*=>', content)
        if home_match and not any(r["path"] == "/" for r in routes):
            routes.append({"path": "/", "name": "Home"})

        # Extract resources (which create index routes)
        resource_matches = re.findall(r'resources\s+:(\w+)', content)
        for resource in resource_matches:
            routes.append({
                "path": f"/{resource}",
                "name": resource.capitalize()
            })

        # Extract custom GET routes
        get_matches = re.findall(r'get\s+"([^"]+)"\s*=>\s*"([^"]+)#([^"]+)"', content)
        for path, controller, action in get_matches:
            if path not in ["/", "home", "up", "service-worker", "manifest"]:
                display_name = path.replace("/", "").replace("_", " ").title() or "Home"
                routes.append({
                    "path": f"/{path}" if not path.startswith("/") else path,
                    "name": display_name
                })

        # Remove duplicates based on path
        seen = set()
        unique_routes = []
        for route in routes:
            if route["path"] not in seen:
                seen.add(route["path"])
                unique_routes.append(route)

        return {"routes": unique_routes}

    except Exception as e:
        print(f"Error parsing routes: {e}")
        return {"routes": [{"path": "/", "name": "Home"}]}


@router.get("/check")
def check_timestamp(request: Request):
    """Returns the timestamp of last message from user in UTC."""
    return {"timestamp": request.app.state.timestamp}


@router.post("/api/update-activity", response_class=JSONResponse)
async def update_activity(request: Request, username: str = Depends(auth)):
    """Update last activity timestamp (called by frontend on user activity)."""
    from datetime import datetime, timezone
    request.app.state.timestamp = datetime.now(timezone.utc)
    return {"timestamp": request.app.state.timestamp.isoformat()}


class FeedbackRequest(BaseModel):
    thread_id: str
    rating: str
    scope: str = "message"
    note: str | None = None
    content: str | None = None
    sent_at: str | None = None
    # Stable id of the rated assistant message, echoed back from the bubble. Lets the
    # mothership join on identity instead of comparing message bodies — text matching
    # silently fabricated a placeholder row whenever a stream was truncated (#688).
    message_key: str | None = None
    # Bounded, redacted browser snapshot (ws close code, reconnect count, recent
    # console output, model/mode). See frontend/chat/utils/LeoDiagnostics.js.
    debug_context: dict | None = None


# A snapshot bigger than this is a bug in the collector, not useful evidence. Drop it
# rather than reject the feedback — losing the user's rating is the worse outcome.
_MAX_DEBUG_CONTEXT_BYTES = 64_000


@router.post("/api/feedback", response_class=JSONResponse)
async def api_submit_feedback(request: Request, body: FeedbackRequest, username: str = Depends(auth)):
    """
    Forward an end-user 👍/👎 to the mothership.

    Thin same-origin passthrough (the browser is already authed to this box),
    mirroring /api/update-activity. Best-effort: a reporting hiccup must never
    500 the browser, so any failure returns {"success": False}.
    """
    if body.rating not in ("good", "bad"):
        return {"success": False, "error": "rating must be 'good' or 'bad'"}

    mothership = getattr(request.app.state, "mothership_client", None)
    if mothership is None:
        from app.services.mothership_client import MothershipClient
        mothership = MothershipClient()
    if not mothership.reporting_enabled:
        return {"success": False, "reason": "mothership_not_configured"}

    debug_context = body.debug_context
    if debug_context is not None:
        try:
            if len(json.dumps(debug_context)) > _MAX_DEBUG_CONTEXT_BYTES:
                logger.warning("feedback debug_context too large; dropping it")
                debug_context = None
        except (TypeError, ValueError):
            debug_context = None

    conn = (debug_context or {}).get("connection") or {}
    last_close = conn.get("last_close") or {}
    logger.info(
        "feedback_submit scope=%s rating=%s thread_id=%s ws_recent_close_code=%s "
        "reconnect_attempts=%s recent_events=%s",
        body.scope, body.rating, body.thread_id,
        last_close.get("code"), conn.get("reconnect_attempts"),
        len((debug_context or {}).get("recent_events") or []),
    )

    from app.services import user_context as user_ctx

    result = await mothership.submit_feedback(
        thread_id=body.thread_id,
        rating=body.rating,
        scope=body.scope,
        note=body.note,
        content=body.content,
        message_key=body.message_key,
        sent_at=body.sent_at,
        debug_context=debug_context,
        # HTTP path: no turn stamp to fall back on, so resolve the signed-in
        # user from the session cookie here.
        user=user_ctx.for_request(request),
    )
    return {"success": bool(result and result.get("success"))}


class FrontendErrorRequest(BaseModel):
    error_class: str
    error_message: str
    stack: str | None = None
    thread_id: str | None = None
    agent_mode: str | None = None
    model: str | None = None
    fingerprint: str | None = None


# A runaway page must not be able to fire huge payloads at us. Generous enough for
# a real stack trace (5000 chars) plus the rest of the envelope.
_MAX_FRONTEND_ERROR_BYTES = 20_000


@router.post("/api/frontend-error", response_class=JSONResponse)
async def api_report_frontend_error(
    request: Request, body: FrontendErrorRequest, username: str = Depends(auth)
):
    """
    Forward a browser-side chat error to the mothership.

    Error telemetry was backend-only: `_report_error_to_mothership` fires on Python
    exceptions, but a client-side socket drop raises nothing server-side. Kody hit
    "Lost connection" mid-run on 2026-07-24 and it left no InstanceError row at all.

    Thin same-origin passthrough (the browser is already authed to this box),
    mirroring /api/feedback. Best-effort: a reporting hiccup must never 500 the
    browser, so any failure returns {"success": False}.
    """
    total = len(body.error_message or "") + len(body.stack or "")
    if total > _MAX_FRONTEND_ERROR_BYTES:
        return {"success": False, "error": "payload too large"}

    mothership = getattr(request.app.state, "mothership_client", None)
    if mothership is None:
        from app.services.mothership_client import MothershipClient
        mothership = MothershipClient()
    if not mothership.reporting_enabled:
        return {"success": False, "reason": "mothership_not_configured"}

    error_message = (body.error_message or "")[:2000]
    stack = (body.stack or "")[:5000]

    fingerprint = body.fingerprint
    if not fingerprint:
        # Same md5 the backend uses (websocket/request_handler.py), so frontend and
        # backend rows for the same failure dedupe together on the mothership.
        import hashlib

        first_line = (error_message.splitlines() or [""])[0]
        fingerprint = hashlib.md5(
            f"{body.error_class}|{first_line[:160]}|{body.agent_mode}".encode("utf-8", "replace")
        ).hexdigest()

    try:
        from datetime import datetime, timezone

        from app.services import user_context as user_ctx

        await mothership.report_error(
            thread_id=body.thread_id,
            error_class=body.error_class,
            error_message=error_message,
            traceback_str=stack,
            agent_mode=body.agent_mode,
            model=body.model,
            llamabot_version=get_container_version(),
            occurred_at=datetime.now(timezone.utc).isoformat(),
            fingerprint=fingerprint,
            source="frontend",
            user=user_ctx.for_request(request),
        )
    except Exception as e:
        logger.warning(f"Frontend error report failed: {e}")
        return {"success": False}

    return {"success": True}


@router.get("/api/lease-status", response_class=JSONResponse)
async def get_lease_status(request: Request):
    """Debug endpoint: show lease manager status."""
    from datetime import datetime, timezone

    mothership = getattr(request.app.state, 'mothership_client', None)
    last_activity = getattr(request.app.state, 'timestamp', None)
    now = datetime.now(timezone.utc)

    if mothership is None:
        return {
            "mothership_enabled": False,
            "error": "Mothership client not initialized"
        }

    return {
        "mothership_enabled": mothership.enabled,
        "instance_name": mothership.instance_name,
        "last_activity": last_activity.isoformat() if last_activity else None,
        "seconds_since_activity": (now - last_activity).total_seconds() if last_activity else None,
        "lease_duration_seconds": mothership.lease_duration_seconds,
    }


@router.get("/api/instance-info", response_class=JSONResponse)
async def get_instance_info():
    """Get instance info from .leonardo/instance.json if it exists."""
    instance_file = ".leonardo/instance.json"

    if not os.path.exists(instance_file):
        return {"instance_name": None}

    try:
        with open(instance_file, "r") as f:
            data = json.load(f)
        return {"instance_name": data.get("instance_name")}
    except Exception as e:
        logger.warning(f"Error reading instance.json: {e}")
        return {"instance_name": None}


@router.post("/api/capture-rails-logs", response_class=JSONResponse)
async def capture_rails_logs_endpoint(request: Request, username: str = Depends(auth)):
    """Capture Rails logs for 10 seconds and return the content."""
    from pathlib import Path
    from app.agents.leonardo.rails_agent.tools import capture_rails_logs

    try:
        file_path = capture_rails_logs(duration=10)
        logs = Path(file_path).read_text()
        return {"logs": logs, "file_path": file_path}
    except Exception as e:
        return JSONResponse(
            status_code=500,
            content={"error": str(e), "detail": f"Failed to capture logs: {str(e)}"}
        )


# ============== Prompt Library API ==============

@router.get("/api/prompts", response_class=JSONResponse)
async def api_get_prompts(
    username: str = Depends(auth),
    session: Session = Depends(get_db_session),
    group: Optional[str] = Query(None, description="Filter by group"),
    search: Optional[str] = Query(None, description="Search term")
):
    """Get all prompts, optionally filtered by group or search term."""
    from app.services.prompt_service import (
        get_all_prompts, get_prompts_by_group, search_prompts
    )

    if search:
        prompts = search_prompts(session, search)
    elif group:
        prompts = get_prompts_by_group(session, group)
    else:
        prompts = get_all_prompts(session)

    return [
        {
            "id": p.id,
            "name": p.name,
            "content": p.content,
            "description": p.description,
            "group": p.group,
            "usage_count": p.usage_count,
            "created_at": p.created_at.isoformat() if p.created_at else None,
            "updated_at": p.updated_at.isoformat() if p.updated_at else None,
        }
        for p in prompts
    ]


@router.get("/api/prompts/groups", response_class=JSONResponse)
async def api_get_prompt_groups(
    username: str = Depends(auth),
    session: Session = Depends(get_db_session)
):
    """Get list of unique prompt groups."""
    from app.services.prompt_service import get_prompt_groups
    groups = get_prompt_groups(session)
    return {"groups": groups}


@router.get("/api/prompts/{prompt_id}", response_class=JSONResponse)
async def api_get_prompt(
    prompt_id: int,
    username: str = Depends(auth),
    session: Session = Depends(get_db_session)
):
    """Get a specific prompt by ID."""
    from app.services.prompt_service import get_prompt_by_id
    prompt = get_prompt_by_id(session, prompt_id)
    if not prompt:
        raise HTTPException(status_code=404, detail="Prompt not found")

    return {
        "id": prompt.id,
        "name": prompt.name,
        "content": prompt.content,
        "description": prompt.description,
        "group": prompt.group,
        "usage_count": prompt.usage_count,
        "created_at": prompt.created_at.isoformat() if prompt.created_at else None,
        "updated_at": prompt.updated_at.isoformat() if prompt.updated_at else None,
    }


@router.post("/api/prompts", response_class=JSONResponse)
async def api_create_prompt(
    request: CreatePromptRequest,
    username: str = Depends(auth),
    session: Session = Depends(get_db_session)
):
    """Create a new prompt."""
    from app.services.prompt_service import create_prompt

    if not request.name or not request.name.strip():
        raise HTTPException(status_code=400, detail="Name is required")
    if not request.content or not request.content.strip():
        raise HTTPException(status_code=400, detail="Content is required")

    prompt = create_prompt(
        session,
        name=request.name,
        content=request.content,
        group=request.group,
        description=request.description
    )

    return {
        "id": prompt.id,
        "name": prompt.name,
        "message": "Prompt created successfully"
    }


@router.patch("/api/prompts/{prompt_id}", response_class=JSONResponse)
async def api_update_prompt(
    prompt_id: int,
    request: UpdatePromptRequest,
    username: str = Depends(auth),
    session: Session = Depends(get_db_session)
):
    """Update an existing prompt."""
    from app.services.prompt_service import update_prompt

    prompt = update_prompt(
        session, prompt_id,
        name=request.name,
        content=request.content,
        group=request.group,
        description=request.description,
        is_active=request.is_active
    )

    if not prompt:
        raise HTTPException(status_code=404, detail="Prompt not found")

    return {"message": "Prompt updated successfully"}


@router.delete("/api/prompts/{prompt_id}", response_class=JSONResponse)
async def api_delete_prompt(
    prompt_id: int,
    username: str = Depends(auth),
    session: Session = Depends(get_db_session),
    hard_delete: bool = Query(False, description="Permanently delete")
):
    """Delete a prompt (soft delete by default)."""
    from app.services.prompt_service import delete_prompt

    if not delete_prompt(session, prompt_id, hard_delete=hard_delete):
        raise HTTPException(status_code=404, detail="Prompt not found")

    return {"message": "Prompt deleted successfully"}


@router.post("/api/prompts/{prompt_id}/use", response_class=JSONResponse)
async def api_use_prompt(
    prompt_id: int,
    username: str = Depends(auth),
    session: Session = Depends(get_db_session)
):
    """Increment usage count when a prompt is attached to a message."""
    from app.services.prompt_service import increment_usage

    prompt = increment_usage(session, prompt_id)
    if not prompt:
        raise HTTPException(status_code=404, detail="Prompt not found")

    return {"usage_count": prompt.usage_count}


# ============== LEONARDO.md API ==============

class UpdateLeonardoMdRequest(BaseModel):
    content: str


@router.get("/api/leonardo-md", response_class=JSONResponse)
async def get_leonardo_md(username: str = Depends(auth)):
    """Get LEONARDO.md content if it exists."""
    leonardo_md_path = ".leonardo/LEONARDO.md"

    if not os.path.exists(leonardo_md_path):
        return {"content": None, "exists": False}

    try:
        with open(leonardo_md_path, "r", encoding="utf-8") as f:
            content = f.read()
        return {"content": content, "exists": True}
    except Exception as e:
        logger.warning(f"Error reading LEONARDO.md: {e}")
        return {"content": None, "exists": False, "error": str(e)}


@router.put("/api/leonardo-md", response_class=JSONResponse)
async def update_leonardo_md(
    request: UpdateLeonardoMdRequest,
    current_user: User = Depends(engineer_or_admin_required)
):
    """Update LEONARDO.md content (engineer/admin only)."""
    leonardo_dir = ".leonardo"
    leonardo_md_path = f"{leonardo_dir}/LEONARDO.md"

    # Ensure .leonardo directory exists
    os.makedirs(leonardo_dir, exist_ok=True)

    try:
        with open(leonardo_md_path, "w", encoding="utf-8") as f:
            f.write(request.content)

        logger.info(f"User '{current_user.username}' updated LEONARDO.md")
        return {"message": "LEONARDO.md updated successfully"}
    except Exception as e:
        logger.error(f"Error writing LEONARDO.md: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to write file: {e}")


# ============== Brand Guide API ==============
#
# The Brand Guide is a small, user-editable branding guideline (named colors,
# logos/icons, free-form notes). Persistence (brand.json + BRAND.md + the
# brand-guidelines skill) lives in app.services.brand_service, shared with the
# agent tools (read_brand_guide / write_brand_guide) so the two write surfaces
# can never drift.

from app.services.brand_service import (
    load_brand,
    save_brand,
    brand_exists,
    DEFAULT_BRAND,
)


class BrandColorModel(BaseModel):
    name: str = ""
    hex: str = ""


class BrandLogoModel(BaseModel):
    name: str = ""
    path: str = ""


class BrandGuideModel(BaseModel):
    colors: list[BrandColorModel] = []
    logos: list[BrandLogoModel] = []
    notes: str = ""


@router.get("/api/brand", response_class=JSONResponse)
async def get_brand(username: str = Depends(auth)):
    """Get the structured brand guide (brand.json), or sensible defaults."""
    return {"brand": load_brand(), "exists": brand_exists()}


@router.put("/api/brand", response_class=JSONResponse)
async def update_brand(request: BrandGuideModel, username: str = Depends(auth)):
    """Persist the brand guide: write brand.json, BRAND.md, and refresh the skill."""
    try:
        brand = save_brand(request.model_dump())
        logger.info(f"User '{username}' updated the brand guide")
        return {"message": "Brand guide saved", "brand": brand}
    except Exception as e:
        logger.error(f"Error writing brand guide: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to write brand guide: {e}")


# ============== Visible Agents Configuration ==============

# Default visible agents for users without a custom configuration
DEFAULT_VISIBLE_AGENTS = ["ticket", "engineer", "testing", "feedback", "user", "beginner"]

# All valid agent mode keys (must match config.js agentModes)
VALID_AGENT_MODES = ["ticket", "engineer", "feedback", "ai_builder", "testing", "user", "beginner"]


class UpdateVisibleAgentsRequest(BaseModel):
    visible_agents: list[str]


@router.get("/api/user/visible-agents", response_class=JSONResponse)
async def get_visible_agents(
    current_user: User = Depends(get_current_user)
):
    """Get current user's visible agents list."""
    if current_user.visible_agents:
        try:
            agents = json.loads(current_user.visible_agents)
            return {"visible_agents": agents}
        except json.JSONDecodeError:
            pass
    return {"visible_agents": DEFAULT_VISIBLE_AGENTS}


@router.put("/api/user/visible-agents", response_class=JSONResponse)
async def set_visible_agents(
    request: UpdateVisibleAgentsRequest,
    current_user: User = Depends(get_current_user),
    session: Session = Depends(get_db_session)
):
    """Set current user's visible agents list."""
    from datetime import datetime, timezone

    # Validate agent keys
    invalid = [a for a in request.visible_agents if a not in VALID_AGENT_MODES]
    if invalid:
        raise HTTPException(status_code=400, detail=f"Invalid agent modes: {invalid}")

    if len(request.visible_agents) == 0:
        raise HTTPException(status_code=400, detail="Must have at least one visible agent")

    current_user.visible_agents = json.dumps(request.visible_agents)
    current_user.updated_at = datetime.now(timezone.utc)
    session.add(current_user)
    session.commit()

    logger.info(f"User '{current_user.username}' updated visible_agents to {request.visible_agents}")
    return {"visible_agents": request.visible_agents, "message": "Visible agents updated"}


# ============== Site Settings API ==============

VALID_SITE_SETTINGS = {
    "show_token_wheel",
    "proactive_build_after_ticket",
    "enable_browser_inspect",
    "enable_live_browser_tools",
    # The VS Code editor: one key owns both the Code tab and the code-server
    # container. Writing it here only moves the tab — POST /api/vscode/enable is
    # the route that also starts or stops the container, so no other setting
    # write can ever reach docker.
    "enable_vscode",
    # Comma-separated tab targets for the chat browser pane. Unlike the others
    # this is not a boolean; ui.resolve_visible_tabs() parses it and drops any
    # name it doesn't recognize, so a bad write degrades to fewer tabs rather
    # than to a broken pane.
    "visible_tabs",
}


# ============== Role → agent-mode permissions (admin only) ==============
# Deliberately NOT part of VALID_SITE_SETTINGS: that PUT is engineer-or-admin,
# and an engineer editing role grants could widen another role's access. These
# endpoints are admin_required, and they validate the payload — the resolver in
# app/permissions.py trusts nothing, but a 400 here beats a silently-ignored key.


@router.get("/api/role-modes", response_class=JSONResponse)
async def api_get_role_modes(
    admin: User = Depends(admin_required),
    session: Session = Depends(get_db_session),
):
    """Current role → agent-mode grants, plus the catalog to render a picker from."""
    from app.permissions import DEFAULT_ROLE_MODES, MODE_AGENTS, custom_modes, get_role_modes

    custom = custom_modes()
    return {
        "role_modes": get_role_modes(session),
        "defaults": DEFAULT_ROLE_MODES,
        "all_modes": sorted(MODE_AGENTS) + sorted(custom),
        "custom_modes": sorted(custom),
    }


@router.put("/api/role-modes", response_class=JSONResponse)
async def api_set_role_modes(
    request: Request,
    admin: User = Depends(admin_required),
    session: Session = Depends(get_db_session),
):
    """Replace the role → agent-mode grants (admin only).

    Body: ``{"role_modes": {"user": ["feedback"], "engineer": [...]}}``.
    Rejects unknown mode keys so a typo fails loudly at the edit rather than
    quietly narrowing someone's access at the next login.
    """
    from datetime import datetime, timezone

    from app.models import SiteSetting
    from app.permissions import ROLE_MODES_SETTING_KEY, custom_modes, known_modes

    body = await request.json()
    role_modes = body.get("role_modes")
    if not isinstance(role_modes, dict):
        raise HTTPException(status_code=400, detail="role_modes must be an object")

    valid = known_modes(custom_modes())
    cleaned: dict[str, list[str]] = {}
    for role, modes in role_modes.items():
        if not isinstance(role, str) or not role.strip():
            raise HTTPException(status_code=400, detail="Role names must be non-empty strings")
        if not isinstance(modes, list) or not all(isinstance(m, str) for m in modes):
            raise HTTPException(status_code=400, detail=f"Modes for '{role}' must be a list of strings")
        unknown = sorted(set(modes) - valid)
        if unknown:
            raise HTTPException(status_code=400, detail=f"Unknown mode(s) for '{role}': {', '.join(unknown)}")
        # De-dupe, preserve the admin's ordering (it drives the dropdown order).
        cleaned[role] = list(dict.fromkeys(modes))

    value = json.dumps(cleaned)
    if len(value) > 1000:  # SiteSetting.value is max_length=1000
        raise HTTPException(status_code=400, detail="Too many roles/modes to store")

    setting = session.get(SiteSetting, ROLE_MODES_SETTING_KEY)
    if setting:
        setting.value = value
        setting.updated_at = datetime.now(timezone.utc)
    else:
        setting = SiteSetting(key=ROLE_MODES_SETTING_KEY, value=value)
        session.add(setting)
    session.commit()

    logger.info(f"Role modes set to {value} by {admin.username}")
    return {"role_modes": cleaned}


def get_site_setting(session: Session, key: str, default: str = "false") -> str:
    """Get a site setting value, returning default if not found.

    Falls back to the default when the auth database is unavailable
    (e.g. LEONARDO_DB_URI is not set), matching db.py's degradation.
    """
    from app.models import SiteSetting
    try:
        setting = session.get(SiteSetting, key)
    except Exception as e:
        logger.warning(f"Could not read site setting '{key}', using default '{default}': {e}")
        return default
    return setting.value if setting else default


@router.get("/api/site-settings/{key}", response_class=JSONResponse)
async def api_get_site_setting(
    key: str,
    username: str = Depends(auth),
    session: Session = Depends(get_db_session),
):
    """Get a site setting value."""
    if key not in VALID_SITE_SETTINGS:
        raise HTTPException(status_code=400, detail=f"Unknown setting: {key}")
    value = get_site_setting(session, key)
    return {"key": key, "value": value}


@router.put("/api/site-settings/{key}", response_class=JSONResponse)
async def api_set_site_setting(
    key: str,
    request: Request,
    current_user: User = Depends(get_current_user),
    session: Session = Depends(get_db_session),
):
    """Set a site setting value (engineer or admin only)."""
    if current_user.role not in ("engineer",) and not current_user.is_admin:
        raise HTTPException(status_code=403, detail="Only engineers or admins can change site settings")
    if key not in VALID_SITE_SETTINGS:
        raise HTTPException(status_code=400, detail=f"Unknown setting: {key}")

    from app.models import SiteSetting
    from datetime import datetime, timezone

    body = await request.json()
    value = str(body.get("value", "false"))

    setting = session.get(SiteSetting, key)
    if setting:
        setting.value = value
        setting.updated_at = datetime.now(timezone.utc)
    else:
        setting = SiteSetting(key=key, value=value)
        session.add(setting)
    session.commit()

    logger.info(f"Site setting '{key}' set to '{value}' by {current_user.username}")
    return {"key": key, "value": value}


# ============== Code editor (VS Code / code-server) ==============
#
# One setting, `enable_vscode`, owns both the Code tab in the chat browser pane
# and the code-server container. These routes are the only place a setting write
# also touches docker. See app/services/vscode_service.py for the two guards
# that keep the editor from starting itself.


def _write_site_setting(session: Session, key: str, value: str) -> None:
    """Persist one site setting. Used by the code editor routes."""
    from datetime import datetime, timezone

    from app.models import SiteSetting

    setting = session.get(SiteSetting, key)
    if setting:
        setting.value = value
        setting.updated_at = datetime.now(timezone.utc)
    else:
        setting = SiteSetting(key=key, value=value)
        session.add(setting)
    session.commit()


def _require_engineer_or_admin(current_user: User) -> None:
    if current_user.role not in ("engineer",) and not current_user.is_admin:
        raise HTTPException(status_code=403, detail="Only engineers or admins can change the code editor")


@router.get("/api/vscode/status", response_class=JSONResponse)
async def api_vscode_status(
    username: str = Depends(auth),
    session: Session = Depends(get_db_session),
):
    """Report whether the editor is switched on and whether it is running."""
    from app.services import vscode_service

    return {
        "enabled": vscode_service.vscode_enabled(session),
        "running": vscode_service.vscode_running(),
    }


@router.post("/api/vscode/enable", response_class=JSONResponse)
async def api_vscode_enable(
    current_user: User = Depends(get_current_user),
    session: Session = Depends(get_db_session),
):
    """Start the editor container, then show the Code tab.

    The setting is written only after the container starts. A failed start
    therefore leaves the tab hidden, because a visible Code tab pointing at a
    stopped editor is the failure this ordering avoids. The failure text (for
    example a missing VSCODE_PASSWORD) comes back in `output` with a 200, so the
    Settings page can show the reason instead of a bare error.
    """
    from app.services import vscode_service

    _require_engineer_or_admin(current_user)

    result = vscode_service.start_vscode()
    if result["ok"]:
        _write_site_setting(session, vscode_service.VSCODE_SETTING_KEY, "true")
        logger.info(f"Code editor enabled by {current_user.username}")

    return {"enabled": bool(result["ok"]), "ok": result["ok"], "output": result["output"]}


@router.post("/api/vscode/disable", response_class=JSONResponse)
async def api_vscode_disable(
    current_user: User = Depends(get_current_user),
    session: Session = Depends(get_db_session),
):
    """Hide the Code tab, then stop the editor container.

    The setting is written first here, and it is written even when the stop
    fails. Hiding the tab is what the user asked for, and a container that
    refuses to stop must not block that request.
    """
    from app.services import vscode_service

    _require_engineer_or_admin(current_user)

    _write_site_setting(session, vscode_service.VSCODE_SETTING_KEY, "false")
    result = vscode_service.stop_vscode()
    logger.info(f"Code editor disabled by {current_user.username}")

    return {"enabled": False, "ok": result["ok"], "output": result["output"]}


# ============== Instance lock ("your free Leo is about to sleep") ==============
#
# The mothership decides; the instance obeys and remembers. The write path is
# authenticated with the SAME shared secret the instance already uses to call the
# mothership (``mothership_api_token`` in .leonardo/instance.json), so no new
# credential has to be provisioned — and, importantly, the instance owner (who is
# an admin on their own box) cannot unlock themselves through the UI.


def _mothership_authorized(request: Request) -> bool:
    """True when the caller presents the instance's mothership bearer token."""
    import secrets

    from app.services.mothership_client import MothershipClient

    config = getattr(getattr(request.app.state, "mothership_client", None), "config", None)
    if not config:
        config = MothershipClient().config
    expected = (config or {}).get("mothership_api_token")
    if not expected:
        return False

    header = request.headers.get("authorization", "")
    scheme, _, token = header.partition(" ")
    if scheme.lower() != "bearer" or not token:
        return False
    return secrets.compare_digest(token, str(expected))


@router.get("/api/instance-lock", response_class=JSONResponse)
async def api_get_instance_lock(
    username: str = Depends(auth),
    session: Session = Depends(get_db_session),
):
    """Lock state for the chat UI to poll. Signed-in users only, read-only."""
    from app.services.instance_lock import get_lock_state

    return get_lock_state(session)


@router.post("/api/instance-lock", response_class=JSONResponse)
async def api_set_instance_lock(
    request: Request,
    session: Session = Depends(get_db_session),
):
    """Lock or unlock this instance (mothership only).

    Body: ``{"locked": true}``, optionally with ``title`` / ``body`` /
    ``upgrade_url`` to override the modal copy without shipping a new image.
    """
    from app.services.instance_lock import set_lock_state

    if not _mothership_authorized(request):
        raise HTTPException(status_code=401, detail="Mothership authorization required")

    body = await request.json()
    if not isinstance(body, dict) or "locked" not in body:
        raise HTTPException(status_code=400, detail="Body must be an object with a 'locked' boolean")

    try:
        state = set_lock_state(session, body)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    # Mirror onto app.state so the WebSocket gate blocks the very next message
    # even if the auth DB goes away afterwards.
    request.app.state.instance_lock = state
    return state


# ============== Environment variables API ==============
#
# There is NO endpoint here that returns the contents of the .env file — no
# values, no key names, not even a "configured" bit. That is the whole design:
# the file holds every provider API key, the database URLs, the SSO login secret
# and the VS Code password, so the browser is never told anything about it.
# Deleting the old reveal endpoint was the point, not an oversight; do not add
# one back.
#
# What IS exposed:
#   * engineer-or-admin — the two allowlisted boolean switches, and the model
#     enabled/disabled list.
#   * admin only — creating and deleting the user's OWN custom variables.
#
# The service layer enforces the same rules independently, so a future route
# cannot widen this by accident.


@router.get("/api/env-vars", response_class=JSONResponse)
async def api_list_env_vars(
    current_user: User = Depends(engineer_or_admin_required),
    session: Session = Depends(get_db_session),
):
    """Settings payload: switches, custom-variable names, restart state.

    Deliberately does NOT include the .env inventory. ``writable`` is a
    capability bit for the UI, not information about the file's contents.
    """
    from app.services import env_settings_service as envsvc
    from app.services import env_store

    return {
        "writable": envsvc.is_writable(),
        "can_edit": bool(current_user.is_admin),
        "toggles": envsvc.toggle_states(),
        "custom_vars": env_store.list_custom_vars(session),
        "pending": env_store.pending_summary(session),
    }


@router.get("/api/env-vars/pending", response_class=JSONResponse)
async def api_pending_env_changes(
    current_user: User = Depends(engineer_or_admin_required),
    session: Session = Depends(get_db_session),
):
    """Edits written to the file but not yet live. Self-clears after a restart."""
    from app.services import env_store
    return env_store.pending_summary(session)


@router.put("/api/env-toggles/{key}", response_class=JSONResponse)
async def api_set_env_toggle(
    key: str,
    request: Request,
    current_user: User = Depends(engineer_or_admin_required),
    session: Session = Depends(get_db_session),
):
    """Flip one of the two allowlisted boolean switches.

    This is the ONLY write path to a platform variable. It is safe to expose
    because ``set_toggle`` refuses any key outside the allowlist and coerces the
    value to the literal "true"/"false" — no caller-supplied text reaches the file,
    and no key outside the list can be touched.
    """
    from app.services import env_settings_service as envsvc
    from app.services import env_store

    body = await request.json()
    try:
        needs_restart = envsvc.set_toggle(key, body.get("value"))
    except envsvc.EnvValidationError as e:
        raise HTTPException(status_code=400, detail=str(e))

    if needs_restart:
        env_store.mark_pending(session, key, "", current_user.username)

    logger.info("Env toggle '%s' set by '%s'", key, current_user.username)
    return {
        "key": key,
        "enabled": envsvc.env_bool(key, False),
        "needs_restart": needs_restart,
        "pending": env_store.pending_summary(session),
    }


@router.post("/api/custom-env-vars", response_class=JSONResponse)
async def api_create_custom_env_var(
    request: Request,
    admin: User = Depends(admin_required),
    session: Session = Depends(get_db_session),
):
    """Define a custom variable, stored in the DB and merged into .env.

    It reaches the Rails container because both services load the same
    ``env_file`` — but only on the next recreate, so this always reports pending.
    """
    from app.services import env_settings_service as envsvc
    from app.services import env_store

    body = await request.json()
    try:
        result = env_store.upsert_custom_var(
            session,
            name=str(body.get("name", "")),
            value=str(body.get("value", "")),
            description=body.get("description"),
            user_id=admin.id,
        )
    except envsvc.EnvValidationError as e:
        raise HTTPException(status_code=400, detail=str(e))

    logger.info("Custom env var '%s' saved by admin '%s'", result["name"], admin.username)
    return {**result, "needs_restart": True, "pending": env_store.pending_summary(session)}


@router.delete("/api/custom-env-vars/{name}", response_class=JSONResponse)
async def api_delete_custom_env_var(
    name: str,
    admin: User = Depends(admin_required),
    session: Session = Depends(get_db_session),
):
    """Remove a custom variable and re-render the managed block."""
    from app.services import env_store

    if not env_store.delete_custom_var(session, name):
        raise HTTPException(status_code=404, detail=f"No custom variable named {name}")

    logger.info("Custom env var '%s' deleted by admin '%s'", name, admin.username)
    return {"name": name, "deleted": True, "pending": env_store.pending_summary(session)}


# ============== Skills Library API (filesystem: .leonardo/skills/<slug>/SKILL.md) ==============
# Skills are now Agent Skills on disk (the SKILL.md open standard), managed the
# same way as LEONARDO.md. The agent authors/uses them via the use_skill /
# write_skill / edit_skill tools; these endpoints back an optional management UI.

@router.get("/api/skills", response_class=JSONResponse)
async def api_get_skills(username: str = Depends(auth)):
    """List all installed skills (slug, name, description)."""
    from app.agents.leonardo.skills import list_all_skills
    return [
        {"slug": s["slug"], "name": s["name"], "description": s["description"]}
        for s in list_all_skills()
    ]


@router.get("/api/skills/{slug}", response_class=JSONResponse)
async def api_get_skill(slug: str, username: str = Depends(auth)):
    """Get a single skill's full SKILL.md (frontmatter + body)."""
    from app.agents.leonardo.skills import get_skill, get_skill_body
    skill = get_skill(slug)
    if not skill:
        raise HTTPException(status_code=404, detail="Skill not found")
    return {
        "slug": skill["slug"],
        "name": skill["name"],
        "description": skill["description"],
        "body": skill["body"],
        "raw": get_skill_body(slug),
    }


@router.put("/api/skills/{slug}", response_class=JSONResponse)
async def api_write_skill(
    slug: str,
    request: WriteSkillRequest,
    current_user: User = Depends(engineer_or_admin_required),
):
    """Create or overwrite a skill (engineer/admin only)."""
    from app.agents.leonardo.skills import write_skill_file
    try:
        saved_slug = write_skill_file(
            request.name, request.description, request.content, slug=request.slug or slug
        )
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    logger.info(f"User '{current_user.username}' wrote skill '{saved_slug}'")
    return {"slug": saved_slug, "message": "Skill saved successfully"}


@router.delete("/api/skills/{slug}", response_class=JSONResponse)
async def api_delete_skill(
    slug: str,
    current_user: User = Depends(engineer_or_admin_required),
):
    """Delete a skill by slug (engineer/admin only)."""
    from app.agents.leonardo.skills import delete_skill_file
    if not delete_skill_file(slug):
        raise HTTPException(status_code=404, detail="Skill not found")
    logger.info(f"User '{current_user.username}' deleted skill '{slug}'")
    return {"message": "Skill deleted successfully"}


# ============== Cookbook API (proxy for llamapress.ai/cookbook.json) ==============
# The published cookbook index sends no CORS headers, so the browser can't fetch it
# directly — this proxies it for the /cookbook slash menu. Cached in memory so
# opening the menu repeatedly doesn't hammer llamapress.ai, and the last good
# payload is served if the fetch fails (a down marketing site must not empty the menu).

COOKBOOK_INDEX_URL = os.getenv("COOKBOOK_URL", "https://llamapress.ai/cookbook.json")
COOKBOOK_SITE_URL = COOKBOOK_INDEX_URL.rsplit("/cookbook.json", 1)[0] or "https://llamapress.ai"
COOKBOOK_CACHE_TTL_SECONDS = 15 * 60

# {"guides": [...], "fetched_at": monotonic seconds}
_cookbook_cache: dict = {"guides": None, "fetched_at": 0.0}


def _normalize_cookbook_guides(payload) -> list:
    """Shape the published index into what the slash menu needs.

    Accepts either {"guides": [...]} or a bare list, and drops entries with no
    slug (nothing to link to). Unknown extra fields are ignored, so the menu keeps
    working when the cookbook grows new ones.
    """
    if isinstance(payload, dict):
        raw = payload.get("guides") or []
    elif isinstance(payload, list):
        raw = payload
    else:
        raw = []

    guides = []
    for entry in raw:
        if not isinstance(entry, dict):
            continue
        slug = str(entry.get("slug") or "").strip()
        if not slug:
            continue
        tags = entry.get("tags") or []
        guides.append({
            "slug": slug,
            "title": str(entry.get("title") or slug),
            "category": str(entry.get("category") or ""),
            "summary": str(entry.get("summary") or ""),
            "tags": [str(t) for t in tags if isinstance(t, (str, int, float))],
            "url": f"{COOKBOOK_SITE_URL}/cookbook/{slug}",
        })
    return guides


async def _fetch_cookbook_index() -> list:
    """Fetch + normalize the published cookbook index. Raises on failure."""
    import httpx

    async with httpx.AsyncClient(timeout=8.0) as client:
        response = await client.get(COOKBOOK_INDEX_URL)
        response.raise_for_status()
        return _normalize_cookbook_guides(response.json())


# The owner's personal list is per-instance-owner and changes the moment they publish, so
# it gets its own short-TTL cache rather than riding the process-global fleet cache.
_personal_cookbook_cache: dict = {}
PERSONAL_COOKBOOK_CACHE_TTL_SECONDS = 60


def _normalize_personal_recipes(payload) -> list:
    """Shape the owner's recipes into the same guide dict the slash menu already consumes.

    Unlisted recipes are kept: they are the owner's own, and hiding them here would mean a
    user could not find a recipe they had just published. Entries without a slug, or a
    payload with no handle, are dropped — there is no resolvable URL for either.
    """
    if not isinstance(payload, dict):
        return []
    handle = str(payload.get("handle") or "").strip()
    raw = payload.get("recipes")
    if not handle or not isinstance(raw, list):
        return []

    guides = []
    for entry in raw:
        if not isinstance(entry, dict):
            continue
        slug = str(entry.get("slug") or "").strip()
        if not slug:
            continue
        guides.append({
            "slug": slug,
            "title": str(entry.get("title") or slug),
            "category": str(entry.get("category") or ""),
            "summary": str(entry.get("summary") or ""),
            # Both .json and .md of this URL exist, so the frontend's existing
            # cookbookJsonUrl mention mechanics work unchanged.
            "url": f"https://llamapress.ai/cookbook/u/{handle}/{slug}",
            "handle": handle,
            "visibility": str(entry.get("visibility") or ""),
            "updated_at": str(entry.get("updated_at") or ""),
            "personal": True,
        })
    return guides


def _merge_personal_cookbook(fleet_guides: list, personal_guides) -> list:
    """Owner's recipes first, then the fleet's.

    A personal recipe SHADOWS a fleet guide with the same slug: if the user has their own
    version of a pattern, that is the one they meant.
    """
    if not personal_guides:
        return fleet_guides
    owned = {g["slug"] for g in personal_guides}
    return list(personal_guides) + [g for g in (fleet_guides or []) if g.get("slug") not in owned]


async def _fetch_personal_cookbook(request) -> list:
    """The owner's recipes, cached briefly. Never raises — an empty list is the floor."""
    import time

    cached = _personal_cookbook_cache.get("guides")
    age = time.monotonic() - _personal_cookbook_cache.get("fetched_at", 0.0)
    if cached is not None and age < PERSONAL_COOKBOOK_CACHE_TTL_SECONDS:
        return cached

    if request is None:
        return []
    mothership = getattr(getattr(request, "app", None), "state", None)
    client = getattr(mothership, "mothership_client", None) if mothership else None
    if client is None:
        return []

    try:
        guides = _normalize_personal_recipes(await client.get_personal_cookbook())
    except Exception as e:  # noqa: BLE001 - the fleet menu must survive this
        logger.warning(f"Could not fetch personal cookbook: {e}")
        return cached or []

    _personal_cookbook_cache["guides"] = guides
    _personal_cookbook_cache["fetched_at"] = time.monotonic()
    return guides


@router.get("/api/cookbook", response_class=JSONResponse)
async def api_get_cookbook(request: Request = None, username: str = Depends(auth)):
    """List cookbook recipes for the /cookbook slash menu — the owner's, then the fleet's."""
    import time

    personal = await _fetch_personal_cookbook(request)

    cached = _cookbook_cache.get("guides")
    age = time.monotonic() - _cookbook_cache.get("fetched_at", 0.0)
    if cached is not None and age < COOKBOOK_CACHE_TTL_SECONDS:
        return {"guides": _merge_personal_cookbook(cached, personal), "stale": False}

    try:
        guides = await _fetch_cookbook_index()
        _cookbook_cache["guides"] = guides
        _cookbook_cache["fetched_at"] = time.monotonic()
        return {"guides": _merge_personal_cookbook(guides, personal), "stale": False}
    except Exception as e:
        logger.warning(f"Could not fetch cookbook index from {COOKBOOK_INDEX_URL}: {e}")
        # Serve the last good payload rather than an empty menu — and the owner's own
        # recipes still show even when the fleet index is unreachable.
        return {"guides": _merge_personal_cookbook(cached or [], personal),
                "stale": True, "error": str(e)}


# ============== File Upload to Assets ==============

UPLOAD_ALLOWED_EXTENSIONS = {
    '.png', '.jpg', '.jpeg', '.gif', '.webp', '.svg', '.ico',
    # Spreadsheets — all Excel variants, incl. macro-enabled and binary
    '.xlsx', '.xls', '.xlsm', '.xlsb', '.xltx', '.xltm', '.csv',
    # Documents
    '.pdf', '.docx',
    # Slideshows — PowerPoint (incl. macros), Keynote, OpenDocument
    '.pptx', '.ppt', '.pptm', '.key', '.odp',
    # Media
    '.mp4', '.webm',
    # Data / text formats
    '.xml', '.json', '.txt', '.md', '.yaml', '.yml', '.html', '.htm',
    # Archives
    '.zip',
}

IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.gif', '.webp', '.svg', '.ico'}

# Types a browser can safely render inline. Deliberately excludes SVG: an SVG can
# embed <script>, so serving it inline is an XSS vector — it downloads instead.
# Everything not listed here (Office docs, slideshows, etc.) also downloads; we do
# not render those server-side (no LibreOffice/conversion — keeps the image small
# and the surface area tiny). The OS opens them in the real app.
INLINE_PREVIEW_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.gif', '.webp', '.pdf', '.ico'}

RAILS_ROOT = "/app/app/rails"
IMAGES_DIR = f"{RAILS_ROOT}/app/assets/images"
IMPORTS_DIR = f"{RAILS_ROOT}/app/imports"


@router.post("/api/upload-to-assets", response_class=JSONResponse)
async def upload_to_assets(
    file: UploadFile = File(...),
    username: str = Depends(auth),
):
    """Upload a file: images go to app/assets/images, everything else to app/imports."""
    import pathlib

    if not file.filename:
        raise HTTPException(status_code=400, detail="No filename provided")

    # Validate extension
    ext = pathlib.Path(file.filename).suffix.lower()
    if ext not in UPLOAD_ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"File type '{ext}' not allowed. Allowed: {', '.join(sorted(UPLOAD_ALLOWED_EXTENSIONS))}"
        )

    # Sanitize filename - keep only safe characters
    safe_filename = re.sub(r'[^\w\-.]', '_', file.filename)

    # Route images to assets/images, everything else to app/imports
    if ext in IMAGE_EXTENSIONS:
        dest_dir = IMAGES_DIR
        relative_path = f"app/assets/images/{safe_filename}"
    else:
        dest_dir = IMPORTS_DIR
        relative_path = f"app/imports/{safe_filename}"

    # Ensure directory exists
    os.makedirs(dest_dir, exist_ok=True)

    # Save file
    dest_path = os.path.join(dest_dir, safe_filename)
    contents = await file.read()
    with open(dest_path, "wb") as f:
        f.write(contents)

    logger.info(f"File uploaded: {relative_path} by {username} ({len(contents)} bytes)")

    return {
        "filename": safe_filename,
        "path": relative_path,
        "size": len(contents),
    }


@router.get("/api/uploaded-files", response_class=JSONResponse)
async def list_uploaded_files(username: str = Depends(auth)):
    """List files in app/assets/images and app/imports."""
    import pathlib

    files = []
    for dir_path, rel_prefix in [(IMAGES_DIR, "app/assets/images"), (IMPORTS_DIR, "app/imports")]:
        d = pathlib.Path(dir_path)
        if not d.exists():
            continue
        for f in d.iterdir():
            if f.is_file() and not f.name.startswith('.'):
                stat = f.stat()
                files.append({
                    "filename": f.name,
                    "path": f"{rel_prefix}/{f.name}",
                    "size": stat.st_size,
                    "uploaded_at": stat.st_mtime,
                    "folder": rel_prefix,
                })

    files.sort(key=lambda x: x["uploaded_at"], reverse=True)
    return {"files": files}


MAX_PREVIEW_BYTES = 50 * 1024 * 1024

PREVIEW_PATH_PREFIXES = {
    "app/assets/images": IMAGES_DIR,
    "app/imports": IMPORTS_DIR,
}


def _resolve_uploaded_file(path: str) -> tuple[str, str]:
    """Map an "app/imports/foo.xlsx" style path onto a real file on disk.

    Returns (absolute_path, filename). Raises HTTPException for anything outside
    the two allowed upload roots, for traversal attempts, for missing files, and
    for files above the preview size cap.
    """
    base_dir = None
    filename = None
    for prefix, candidate_base in PREVIEW_PATH_PREFIXES.items():
        if path.startswith(prefix + "/"):
            filename = path[len(prefix) + 1:]
            base_dir = candidate_base
            break
    if base_dir is None or not filename:
        raise HTTPException(status_code=400, detail="Invalid path")

    if "/" in filename or "\\" in filename or filename.startswith("."):
        raise HTTPException(status_code=400, detail="Invalid filename")

    real_full = os.path.realpath(os.path.join(base_dir, filename))
    real_base = os.path.realpath(base_dir)
    if not (real_full == real_base or real_full.startswith(real_base + os.sep)):
        raise HTTPException(status_code=400, detail="Path traversal blocked")

    if not os.path.isfile(real_full):
        raise HTTPException(status_code=404, detail="File not found")

    size = os.path.getsize(real_full)
    if size > MAX_PREVIEW_BYTES:
        raise HTTPException(status_code=413, detail=f"File too large to preview ({size} bytes, max {MAX_PREVIEW_BYTES})")

    return real_full, filename


@router.get("/api/uploaded-files/preview")
async def preview_uploaded_file(path: str, download: bool = False, username: str = Depends(auth)):
    """Serve an uploaded file.

    download=1 always forces a download (Content-Disposition: attachment).
    Otherwise we serve inline only for browser-native types (raster images, PDF);
    everything else — Office docs, slideshows, SVG — downloads. Capped at 50MB.
    """
    import pathlib
    from fastapi.responses import FileResponse

    real_full, filename = _resolve_uploaded_file(path)

    ext = pathlib.Path(filename).suffix.lower()
    inline = (not download) and ext in INLINE_PREVIEW_EXTENSIONS
    disposition = "inline" if inline else "attachment"
    return FileResponse(real_full, filename=filename, content_disposition_type=disposition)


# ---------------------------------------------------------------------------
# Spreadsheet preview — parsed HERE, on the box, never by a third party.
#
# The asset library used to hand spreadsheets to Microsoft's Office Online
# viewer (view.officeapps.live.com) in an iframe. That can never work: the
# viewer fetches the file from Microsoft's servers, and /api/uploaded-files/preview
# requires this user's session. So we do the boring thing — parse the workbook
# with openpyxl (already a dependency, used by the pyxl agent) and hand the
# frontend plain rows of cell values to draw as a grid.
# ---------------------------------------------------------------------------

# openpyxl reads the OOXML family only. Legacy binary .xls/.xlsb are not covered.
SHEET_PREVIEW_EXTENSIONS = {'.xlsx', '.xlsm', '.xltx', '.xltm'}
DELIMITED_PREVIEW_EXTENSIONS = {'.csv': ',', '.tsv': '\t'}

SHEET_PREVIEW_MAX_ROWS = 500
SHEET_PREVIEW_MAX_COLS = 100


def _cell_to_str(value) -> str:
    """Render one cell for display. Dates as ISO, floats without trailing .0."""
    import datetime as _dt

    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value.isoformat(sep=" ") if isinstance(value, _dt.datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


@router.get("/api/uploaded-files/sheet", response_class=JSONResponse)
async def preview_uploaded_sheet(path: str, sheet: int = 0, username: str = Depends(auth)):
    """Parse an uploaded spreadsheet/CSV and return one sheet as rows of strings.

    Response: {"sheets": [names...], "active": i, "rows": [[cell, ...], ...],
               "total_rows": n, "truncated": bool}
    415 means "this format can't be parsed here" — the caller falls back to its
    client-side reader (legacy .xls/.xlsb).
    """
    import csv
    import io
    import pathlib

    real_full, filename = _resolve_uploaded_file(path)
    ext = pathlib.Path(filename).suffix.lower()

    if ext in DELIMITED_PREVIEW_EXTENSIONS:
        delimiter = DELIMITED_PREVIEW_EXTENSIONS[ext]
        with open(real_full, "r", encoding="utf-8-sig", errors="replace", newline="") as fh:
            all_rows = list(csv.reader(io.StringIO(fh.read()), delimiter=delimiter))
        total_rows = len(all_rows)
        rows = [[_cell_to_str(c) for c in r[:SHEET_PREVIEW_MAX_COLS]] for r in all_rows[:SHEET_PREVIEW_MAX_ROWS]]
        return {
            "sheets": [filename],
            "active": 0,
            "rows": rows,
            "total_rows": total_rows,
            "truncated": total_rows > SHEET_PREVIEW_MAX_ROWS,
        }

    if ext not in SHEET_PREVIEW_EXTENSIONS:
        raise HTTPException(status_code=415, detail=f"No server-side preview for {ext} files")

    try:
        import openpyxl
        wb = openpyxl.load_workbook(real_full, data_only=True, read_only=True)
    except HTTPException:
        raise
    except Exception as e:
        logger.warning(f"Spreadsheet preview failed for {filename}: {e}")
        raise HTTPException(status_code=422, detail="Could not read this workbook")

    try:
        names = list(wb.sheetnames)
        if not names:
            raise HTTPException(status_code=422, detail="Workbook has no sheets")
        index = sheet if 0 <= sheet < len(names) else 0
        ws = wb[names[index]]

        rows = []
        total_rows = 0
        for row in ws.iter_rows(values_only=True):
            total_rows += 1
            if len(rows) < SHEET_PREVIEW_MAX_ROWS:
                rows.append([_cell_to_str(c) for c in row[:SHEET_PREVIEW_MAX_COLS]])

        # Trim trailing all-blank rows openpyxl reports for formatted-but-empty cells.
        while rows and not any(c for c in rows[-1]):
            rows.pop()
            total_rows -= 1

        return {
            "sheets": names,
            "active": index,
            "rows": rows,
            "total_rows": total_rows,
            "truncated": total_rows > len(rows),
        }
    finally:
        wb.close()


# ============== Remote Setup API ==============
# Used by the mothership (Rails) to push context files after claiming an instance.
# Auth: HTTP Basic Auth with the admin credentials set up by LlamabotAdminRegistrar.


class WriteFileEntry(BaseModel):
    path: str
    content: str


class WriteFilesRequest(BaseModel):
    files: list[WriteFileEntry]


# Map friendly mothership paths to actual filesystem paths.
# Paths not in this map are written as-is (relative to CWD).
WRITE_FILES_PATH_MAP = {
    "Leonardo.md": ".leonardo/LEONARDO.md",
    "User.md": ".leonardo/USER.md",
}


@router.post("/api/write_files", response_class=JSONResponse)
async def api_write_files(
    body: WriteFilesRequest,
    current_user: User = Depends(admin_required),
):
    """Write files to disk. Used by the mothership to push context files after claiming an instance."""
    results = []

    for entry in body.files:
        file_path = entry.path

        # Resolve friendly names to actual paths
        resolved = WRITE_FILES_PATH_MAP.get(file_path, file_path)
        resolved = os.path.normpath(resolved)

        # Prevent path traversal outside the working directory
        if resolved.startswith("..") or os.path.isabs(resolved):
            results.append({"path": file_path, "status": "error", "detail": "Absolute or traversal paths not allowed"})
            continue

        try:
            parent = os.path.dirname(resolved)
            if parent:
                os.makedirs(parent, exist_ok=True)
            with open(resolved, "w", encoding="utf-8") as f:
                f.write(entry.content)
            logger.info(f"write_files: wrote {resolved} ({len(entry.content)} chars) by {current_user.username}")
            results.append({"path": file_path, "status": "ok"})
        except Exception as e:
            logger.error(f"write_files: failed to write {file_path}: {e}")
            results.append({"path": file_path, "status": "error", "detail": str(e)})

    return {"results": results}


class ImportFromS3Request(BaseModel):
    url: str
    filename: str | None = None


@router.post("/api/setup/import-excel", response_class=JSONResponse)
async def api_import_excel_from_s3(
    body: ImportFromS3Request,
    current_user: User = Depends(admin_required),
):
    """Download a file from S3 (pre-signed URL) and save it to rails/app/imports/."""
    import pathlib
    import httpx

    # Derive filename from the URL if not provided (strip query params)
    if body.filename:
        filename = body.filename
    else:
        url_path = body.url.split("?")[0]
        filename = url_path.rsplit("/", 1)[-1]

    if not filename:
        raise HTTPException(status_code=400, detail="Could not determine filename from URL. Provide a 'filename' field.")

    # Validate extension
    ext = pathlib.Path(filename).suffix.lower()
    if ext not in {".xlsx", ".xls", ".csv"}:
        raise HTTPException(status_code=400, detail=f"File type '{ext}' not allowed. Allowed: .xlsx, .xls, .csv")

    # Sanitize filename
    safe_filename = re.sub(r'[^\w\-.]', '_', filename)

    os.makedirs(IMPORTS_DIR, exist_ok=True)
    dest_path = os.path.join(IMPORTS_DIR, safe_filename)

    try:
        async with httpx.AsyncClient(timeout=60.0) as client:
            resp = await client.get(body.url)
            resp.raise_for_status()

        with open(dest_path, "wb") as f:
            f.write(resp.content)

        logger.info(f"import-excel: downloaded {safe_filename} ({len(resp.content)} bytes) by {current_user.username}")
        return {
            "filename": safe_filename,
            "path": f"app/imports/{safe_filename}",
            "size": len(resp.content),
        }
    except httpx.HTTPStatusError as e:
        raise HTTPException(status_code=502, detail=f"Failed to download from S3: HTTP {e.response.status_code}")
    except Exception as e:
        logger.error(f"import-excel: failed to download {body.url}: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to download file: {e}")


ALLOWED_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
MAX_IMAGE_SIZE = 25 * 1024 * 1024  # 25 MB


@router.post("/api/setup/import-image", response_class=JSONResponse)
async def api_import_image_from_s3(
    body: ImportFromS3Request,
    current_user: User = Depends(admin_required),
):
    """Download an image from S3 and save it to rails/app/imports/."""
    import pathlib
    import httpx
    from urllib.parse import unquote

    # Derive filename
    if body.filename:
        filename = body.filename
    else:
        url_path = body.url.split("?")[0]
        filename = unquote(url_path.rsplit("/", 1)[-1])

    if not filename:
        raise HTTPException(status_code=400, detail="Could not determine filename from URL. Provide a 'filename' field.")

    # Validate extension
    ext = pathlib.Path(filename).suffix.lower()
    if ext not in ALLOWED_IMAGE_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"not an image: {ext}")

    # Sanitize (strip path traversal, unsafe chars)
    safe_filename = re.sub(r'[^\w\-.]', '_', os.path.basename(filename))

    os.makedirs(IMPORTS_DIR, exist_ok=True)
    dest_path = os.path.join(IMPORTS_DIR, safe_filename)

    try:
        async with httpx.AsyncClient(timeout=60.0) as client:
            resp = await client.get(body.url)
            resp.raise_for_status()

        # Size check
        if len(resp.content) > MAX_IMAGE_SIZE:
            raise HTTPException(status_code=400, detail=f"Image too large: {len(resp.content)} bytes (max {MAX_IMAGE_SIZE})")

        # Magic-byte validation
        header = resp.content[:12]
        is_valid_image = (
            header[:8] == b'\x89PNG\r\n\x1a\n'
            or header[:2] == b'\xff\xd8'
            or header[:4] == b'GIF8'
            or (header[:4] == b'RIFF' and header[8:12] == b'WEBP')
        )
        if not is_valid_image:
            raise HTTPException(status_code=400, detail="File content does not match a supported image format")

        with open(dest_path, "wb") as f:
            f.write(resp.content)

        logger.info(f"import-image: downloaded {safe_filename} ({len(resp.content)} bytes) by {current_user.username}")
        return {"status": "ok", "saved_path": f"rails/app/imports/{safe_filename}"}

    except httpx.HTTPStatusError as e:
        raise HTTPException(status_code=502, detail=f"Failed to download from S3: HTTP {e.response.status_code}")
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"import-image: failed to download {body.url}: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to download file: {e}")
